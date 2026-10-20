import openpyxl, re, os, datetime
from copy import copy
W=r"C:\Users\ACE\GREEN GOLD ENERGY Dropbox\Projects\Project (EPC)\2.SA"
src=W+r"\GG-31 Warnertown BESS\Design\Engineering\8. Deliverables\Transmittal\GG-31-TSMT-11.xlsx"
out=W+r"\Waterloo\Design\Engineering\8. Deliverables\Transmittal\GG38-TSMT-01.xlsx"
sub=r"C:\Users\jilin\AppData\Local\Temp\claude\D--\7f4f3d51-b932-4dbf-b86f-b81a2200ce89\scratchpad\GG38_IFR_submission"
files=sorted(f for f in os.listdir(sub) if f.lower().endswith(".pdf"))
wb=openpyxl.load_workbook(src); ws=wb.active
proto={r:[copy(ws.cell(r,c)._style) for c in range(1,8)] for r in (14,15)}
h14,h15=ws.row_dimensions[14].height,ws.row_dimensions[15].height
for r in range(14,ws.max_row+30):
    for c in range(1,8): ws.cell(r,c).value=None
ws["E4"]="No: GG38-TSMT-01"; ws["E5"]="Date: "+datetime.date.today().strftime("%d/%m/%Y")
ws["C6"]=None; ws["C8"]=None
ws["C9"]="GG38 Waterloo BESS, SA"; ws["C10"]="ISSUED FOR REVIEW (IFR) - DRAWINGS & REPORTS"
pat=re.compile(r"^(GG38-[A-Z]-[A-Z]{3}-\d{3}(?:-\d)?)_(.+)_Rev([A-Z0-9]+)\.pdf$")
for i,f in enumerate(files):
    m=pat.match(f); assert m, f
    r=14+2*i
    for rr,key in ((r,14),(r+1,15)):
        for c in range(1,8): ws.cell(rr,c)._style=copy(proto[key][c-1])
    ws.row_dimensions[r].height=h14; ws.row_dimensions[r+1].height=h15
    ws.cell(r,2,i+1); ws.cell(r,3,m.group(1)); ws.cell(r,4,1); ws.cell(r,5,f)
    ws.cell(r,6,"New"); ws.cell(r,7,m.group(2))
wb.save(out); print(len(files),"rows ->",out)
